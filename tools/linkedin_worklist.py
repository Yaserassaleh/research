#!/usr/bin/env python3
"""Manual LinkedIn review worklist: build it from the institution files, and apply the reviewer's answers back.

  python3 tools/linkedin_worklist.py build
      Writes exports/linkedin_review_worklist.xlsx (sheets: Read me, Decision-makers, Board and others) and
      exports/linkedin_review_worklist.csv. One row per person without a verified LinkedIn URL (departed people
      excluded), institutions in ARTA-score order, chief executives and chairs first. The yellow columns are for the
      reviewer; everything else comes from research/institutions/*.json.

  python3 tools/linkedin_worklist.py apply exports/linkedin_review_worklist.xlsx
      Reads the reviewer's columns. For each row marked Confirmed with a profile URL: sets linkedin_url and
      linkedin_status "Verified (profile/search match)", adds a tier-4 source ("LinkedIn profile (manual review)")
      and its id to the person's source_ids. Rows marked Not found: linkedin_status "Not checked" becomes
      "Not verified". Rows marked Ambiguous with a note: the note is added to the institution's open_items.
      Every change is logged in research/central_edits.json. Run tools/build_db.py afterwards.

Why manual: LinkedIn answers HTTP 999 to automated reads and the search tool declines name-based look-ups, so the
automated pass (PEOPLE_DEEP_DIVE.md §7) found 7 URLs in 149 searches. A person can confirm a profile in a minute.
Rules for the reviewer are on the Read me sheet (same as protocol §2 and PEOPLE_DEEP_DIVE.md §7.3-7.4).
"""
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from people_pass import is_decision_maker, load_all  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INST_DIR = os.path.join(ROOT, "research", "institutions")
XLSX = os.path.join(ROOT, "exports", "linkedin_review_worklist.xlsx")
CSV = os.path.join(ROOT, "exports", "linkedin_review_worklist.csv")
CENTRAL = os.path.join(ROOT, "research", "central_edits.json")
TODAY = dt.date.today().isoformat()

DATA_COLS = ["Priority", "Institution", "Institution id", "Country", "ARTA score", "ARTA tier", "Name",
             "Arabic name", "Current title", "Functions", "Entity", "Appointed", "Role status",
             "LinkedIn status (now)", "Education", "Qualifications", "Career (from the database)",
             "Leads already seen (unverified)", "Institution LinkedIn page", "Suggested search"]
REVIEW_COLS = ["Confirmed LinkedIn URL", "Confirmation", "Reviewer", "Date checked", "Notes"]
CONFIRMATION = ["Confirmed", "Not found", "Ambiguous"]
FUNC_RANK = {"Group CEO": 0, "CEO": 0, "Deputy CEO": 1, "Chair": 2, "Wealth": 3, "Private banking": 3,
             "Affluent": 3, "Investments / asset management": 4, "Brokerage": 4, "Digital": 5, "Technology": 5,
             "Innovation / partnerships": 5, "Subsidiary CEO": 6, "Retail": 6, "Strategy": 6}
HONORIFICS = {"sheikh", "sheikha", "dr", "dr.", "h.e.", "he", "eng", "eng.", "mr", "mrs", "ms", "bin", "bint",
              "al", "al-", "el", "abu", "abdul", "abdullah", "mohammed", "mohammad", "ahmed", "ali", "khalid"}
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F3864")
LIGHT = PatternFill("solid", fgColor="F2F2F2")


def name_tokens(name):
    toks = [t.strip(".,()").lower() for t in re.split(r"[\s/]+", name or "")]
    return [t for t in toks if len(t) >= 4 and t not in HONORIFICS]


def leads_for(person, open_items):
    """open_items lines that mention LinkedIn and this person (by a distinctive name token)."""
    toks = name_tokens(person["name"])
    out = []
    for item in open_items or []:
        if "linkedin" not in item.lower():
            continue
        low = item.lower()
        if any(t in low for t in toks):
            out.append(item.strip())
    return " | ".join(out)[:600]


def rows():
    insts = load_all()
    db_path = os.path.join(ROOT, "exports", "database.json")
    meta = {}
    if os.path.exists(db_path):
        for i in json.load(open(db_path, encoding="utf-8"))["institutions"]:
            meta[i["id"]] = i
    order = sorted(insts, key=lambda i: -(meta.get(i, {}).get("score", {}).get("total") or 0))
    dm_rows, other_rows = [], []
    for iid in order:
        d = insts[iid]
        m = meta.get(iid, {})
        people = [p for p in d["people"]
                  if p.get("role_status") != "Departed"
                  and not (p.get("linkedin_status") or "").startswith("Verified")]
        people.sort(key=lambda p: min([FUNC_RANK.get(f, 7) for f in (p.get("functions") or [])] or [7]))
        for p in people:
            search = f'"{p["name"]}" {d.get("short_name") or d["name"]} LinkedIn'
            r = OrderedDict([
                ("Priority", None), ("Institution", d["name"]), ("Institution id", iid),
                ("Country", m.get("country_name") or d.get("country")),
                ("ARTA score", (m.get("score") or {}).get("total")), ("ARTA tier", (m.get("score") or {}).get("tier")),
                ("Name", p["name"]), ("Arabic name", p.get("name_ar")), ("Current title", p.get("title")),
                ("Functions", "; ".join(p.get("functions") or [])), ("Entity", p.get("entity")),
                ("Appointed", p.get("appointed")), ("Role status", p.get("role_status")),
                ("LinkedIn status (now)", p.get("linkedin_status") or "Not checked"),
                ("Education", p.get("education")), ("Qualifications", p.get("qualifications")),
                ("Career (from the database)", (p.get("previous_roles") or "")[:500]),
                ("Leads already seen (unverified)", leads_for(p, d.get("open_items"))),
                ("Institution LinkedIn page", d.get("linkedin_company")), ("Suggested search", search),
            ])
            (dm_rows if is_decision_maker(p) else other_rows).append(r)
    for n, r in enumerate(dm_rows, 1):
        r["Priority"] = n
    for n, r in enumerate(other_rows, 1):
        r["Priority"] = n
    return dm_rows, other_rows


def write_sheet(wb, title, data):
    ws = wb.create_sheet(title)
    cols = DATA_COLS + REVIEW_COLS
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in data:
        ws.append([r[c] for c in DATA_COLS] + [None] * len(REVIEW_COLS))
    n = len(data) + 1
    first_review = len(DATA_COLS) + 1
    for row in ws.iter_rows(min_row=2, max_row=max(n, 2)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column >= first_review:
                cell.fill = YELLOW
    widths = {"Priority": 7, "Institution": 26, "Institution id": 14, "Country": 12, "ARTA score": 8,
              "ARTA tier": 8, "Name": 28, "Arabic name": 22, "Current title": 34, "Functions": 20, "Entity": 24,
              "Appointed": 11, "Role status": 16, "LinkedIn status (now)": 14, "Education": 40,
              "Qualifications": 22, "Career (from the database)": 60, "Leads already seen (unverified)": 50,
              "Institution LinkedIn page": 30, "Suggested search": 36, "Confirmed LinkedIn URL": 44,
              "Confirmation": 14, "Reviewer": 14, "Date checked": 12, "Notes": 44}
    for c, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 16)
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(n, 2)}"
    if n >= 2:
        col = get_column_letter(first_review + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(CONFIRMATION) + '"', allow_blank=True)
        dv.error = "Choose Confirmed, Not found or Ambiguous"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    return ws


def write_readme(wb, n_dm, n_other):
    ws = wb.active
    ws.title = "Read me"
    url_col = get_column_letter(len(DATA_COLS) + 1)
    lines = [
        ("LinkedIn review worklist — ARTA GCC FI Intelligence", True),
        (f"Generated {TODAY} from research/institutions/*.json by tools/linkedin_worklist.py build. "
         "One row per person who has no verified LinkedIn URL (people marked Departed are left out). "
         "Institutions are in ARTA-score order; within an institution, chief executives and chairs come first.", False),
        ("", False),
        ("Rows to review", True),
        ("Decision-makers sheet", False),
        ("Board and others sheet", False),
        ("URLs confirmed so far (both sheets)", False),
        ("", False),
        ("What to edit", True),
        (f"Only the yellow columns {url_col} to {get_column_letter(len(DATA_COLS) + len(REVIEW_COLS))} on the two "
         "sheets: Confirmed LinkedIn URL, Confirmation (Confirmed / Not found / Ambiguous), Reviewer, Date checked, "
         "Notes. The white columns are the database and are overwritten on the next build.", False),
        ("", False),
        ("Rules (RESEARCH_PROTOCOL.md §2; PEOPLE_DEEP_DIVE.md §7.3–7.4)", True),
        ("1. Open the profile. Mark Confirmed only when it shows all three: the same person's name (transliteration "
         "variants are fine: Al/Al-/El, Mohammed/Mohammad, hyphen/space, initials, honorifics), the institution or its "
         "profiled subsidiary as the CURRENT employer, and a title in the same function as the Current title column. "
         "An older title at the same institution still counts if no other employer is shown as current.", False),
        ("2. A similar name, a different current employer, or a title in another function is NOT a match. If the "
         "profile shows the person has moved on or holds a later title, choose Ambiguous and write what it shows in "
         "Notes (it becomes an open item for the research team; the role status is not changed from a profile alone).",
         False),
        ("3. Record the URL as the browser shows it, without anything after a question mark "
         "(https://www.linkedin.com/in/<slug>/ or a country subdomain such as ae.linkedin.com). Never guess or build a URL, "
         "and do not take URLs from data brokers (RocketReach, ZoomInfo, SignalHire, TheOfficialBoard).", False),
        ("4. Professional information only. Copy nothing else from the profile: no contact details, photos, personal "
         "posts, ages or anything outside the person's work history.", False),
        ("5. The 'Leads already seen' column lists LinkedIn snippets the research agents saw but could not verify; "
         "treat them as hints, not answers. 'Suggested search' is the search string a reviewer can type.", False),
        ("", False),
        ("Sending the answers back", True),
        ("Save this file in place (exports/linkedin_review_worklist.xlsx) and run, from the project folder: "
         "python3 tools/linkedin_worklist.py apply exports/linkedin_review_worklist.xlsx  — then python3 tools/build_db.py, "
         "tools/build_workbook.py and tools/build_site.py. Confirmed rows get the URL, status 'Verified (profile/search "
         "match)' and a tier-4 source; Not found rows become 'Not verified'; Ambiguous notes go to open_items. The edits are "
         "logged in research/central_edits.json.", False),
        ("", False),
        ("Example of a completed row (reviewer columns only)", True),
    ]
    r = 1
    for text, flag in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bool(flag is True), size=12 if r == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    # counts (formulas, so they follow the sheets)
    ws.cell(row=5, column=2, value=f"=COUNTA('Decision-makers'!G2:G{max(n_dm + 1, 2)})")
    ws.cell(row=6, column=2, value=f"=COUNTA('Board and others'!G2:G{max(n_other + 1, 2)})")
    ws.cell(row=7, column=2, value=(f"=COUNTA('Decision-makers'!{url_col}2:{url_col}{max(n_dm + 1, 2)})"
                                    f"+COUNTA('Board and others'!{url_col}2:{url_col}{max(n_other + 1, 2)})"))
    for rr in (5, 6, 7):
        ws.cell(row=rr, column=2).font = Font(name=FONT, size=10)
        ws.cell(row=rr, column=1).font = Font(name=FONT, size=10)
    ex_head = ["Name", "Current title"] + REVIEW_COLS
    ex_row = ["David O'Loan", "Deputy Group CEO Treasury and Financial Institutions",
              "https://kw.linkedin.com/in/david-o-loan", "Confirmed", "A. Reviewer", TODAY,
              "Profile: Deputy Group CEO Treasury & FI at Kuwait Finance House (current)"]
    for c, v in enumerate(ex_head, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = LIGHT
    for c, v in enumerate(ex_row, 1):
        cell = ws.cell(row=r + 1, column=c, value=v)
        cell.font = Font(name=FONT, size=10)
        if c >= 3:
            cell.fill = YELLOW
    ws.cell(row=r + 2, column=1, value="(The example is the one verified URL in the database that this list is "
                                        "modelled on; it is not a row to review.)").font = Font(name=FONT, italic=True,
                                                                                               size=9)
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 44
    for c in "CDEFG":
        ws.column_dimensions[c].width = 22


def build():
    dm, other = rows()
    wb = Workbook()
    write_readme(wb, len(dm), len(other))
    write_sheet(wb, "Decision-makers", dm)
    write_sheet(wb, "Board and others", other)
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    with open(CSV, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["Sheet"] + DATA_COLS + REVIEW_COLS)
        for name, data in (("Decision-makers", dm), ("Board and others", other)):
            for r in data:
                w.writerow([name] + [r[c] for c in DATA_COLS] + [""] * len(REVIEW_COLS))
    insts = sum(1 for _ in {r["Institution id"] for r in dm + other})
    print(f"{len(dm)} decision-makers and {len(other)} board/other members without a verified LinkedIn URL "
          f"({insts} institutions) → {os.path.relpath(XLSX, ROOT)}, {os.path.relpath(CSV, ROOT)}")


def clean_url(u):
    u = (u or "").strip()
    if not u:
        return None
    u = u.split("?")[0].split("#")[0]
    if not re.match(r"^https://([a-z]{2,3}\.)?linkedin\.com/in/[^/\s]+/?$", u):
        return None
    return u


def dump_like(path, data):
    """Write JSON in the same style as the existing file (indent, ASCII escaping, trailing newline)."""
    raw = open(path, encoding="utf-8").read()
    m = re.search(r'\n([ \t]+)"', raw)
    txt = json.dumps(data, ensure_ascii=all(ord(c) < 128 for c in raw), indent=m.group(1) if m else 1)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(txt + ("\n" if raw.endswith("\n") else ""))


def log_entry(iid, name, field, change, reviewer, date):
    return {"date": TODAY, "institution_id": iid, "field": f"people[{name}].{field}", "change": change,
            "reason": f"Manual LinkedIn review by {reviewer} on {date} (tools/linkedin_worklist.py apply)",
            "by": "Central QA"}


def next_source_id(sources):
    nums = [int(m.group(1)) for s in sources for m in [re.match(r"S(\d+)$", s.get("id") or "")] if m]
    return f"S{(max(nums) if nums else 0) + 1}"


def apply(path):
    wb = load_workbook(path, data_only=True)
    insts = load_all()
    changed, log, skipped = {}, [], []
    col = {name: i for i, name in enumerate(DATA_COLS + REVIEW_COLS)}
    for sheet in ("Decision-makers", "Board and others"):
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not row or not row[col["Name"]]:
                continue
            iid, name = row[col["Institution id"]], row[col["Name"]]
            verdict = (row[col["Confirmation"]] or "").strip()
            url = clean_url(row[col["Confirmed LinkedIn URL"]])
            reviewer = (row[col["Reviewer"]] or "").strip() or "manual review"
            date = row[col["Date checked"]]
            date = date.date().isoformat() if hasattr(date, "date") else (str(date).strip() if date else TODAY)
            note = (row[col["Notes"]] or "").strip()
            if not verdict and not url:
                continue
            d = insts.get(iid)
            person = next((p for p in (d or {}).get("people", []) if p["name"] == name), None)
            if not person:
                skipped.append(f"{iid} / {name}: not found in the institution file")
                continue
            if verdict == "Confirmed":
                if not url:
                    skipped.append(f"{iid} / {name}: Confirmed but the URL is missing or not a linkedin.com/in/ profile URL")
                    continue
                if person.get("linkedin_url") == url and (person.get("linkedin_status") or "").startswith("Verified"):
                    continue
                sid = next_source_id(d["sources"])
                d["sources"].append({"id": sid, "title": f"{name} - LinkedIn", "publisher": "LinkedIn", "url": url,
                                     "doc_type": "LinkedIn profile (manual review)", "tier": 4, "published": None,
                                     "accessed": date, "language": "en",
                                     "note": f"Profile opened and confirmed by {reviewer} on {date}: same name, "
                                             f"current employer and title as the official record."
                                             + (f" Reviewer note: {note}" if note else "")})
                person["linkedin_url"] = url
                person["linkedin_status"] = "Verified (profile/search match)"
                person.setdefault("source_ids", [])
                if sid not in person["source_ids"]:
                    person["source_ids"].append(sid)
                changed[iid] = d
                log.append(log_entry(iid, name, "linkedin_url", f"Set to {url}; status Verified (profile/search match); "
                                     f"source {sid} added", reviewer, date))
            elif verdict == "Not found":
                if (person.get("linkedin_status") or "Not checked") == "Not checked":
                    person["linkedin_status"] = "Not verified"
                    changed[iid] = d
                    log.append(log_entry(iid, name, "linkedin_status", "Not checked -> Not verified (no matching profile)",
                                         reviewer, date))
            elif verdict == "Ambiguous":
                if note:
                    line = (f"Manual LinkedIn review {date} ({reviewer}): {name} — {note}; verify against official "
                            f"sources before changing the record.")
                    d.setdefault("open_items", [])
                    if line not in d["open_items"]:
                        d["open_items"].append(line)
                        changed[iid] = d
                        log.append(log_entry(iid, name, "open_items", "Line added: LinkedIn profile ambiguous", reviewer, date))
            elif verdict:
                skipped.append(f"{iid} / {name}: unknown Confirmation value '{verdict}'")
    for iid, d in changed.items():
        dump_like(os.path.join(INST_DIR, f"{iid}.json"), d)
    if log:
        central = json.load(open(CENTRAL, encoding="utf-8")) if os.path.exists(CENTRAL) else []
        if isinstance(central, dict):
            central.setdefault("edits", []).extend(log)
        else:
            central.extend(log)
        with open(CENTRAL, "w", encoding="utf-8") as fh:
            json.dump(central, fh, ensure_ascii=False, indent=1)
    print(f"{len(log)} edits in {len(changed)} institution files; {len(skipped)} rows skipped")
    for s in skipped:
        print("  skipped:", s)
    if changed:
        print("Now run: python3 tools/build_db.py && python3 tools/build_workbook.py && python3 tools/build_site.py")


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else "build"
    if cmd == "build":
        build()
    elif cmd == "apply":
        if len(sys.argv) < 3:
            sys.exit("usage: linkedin_worklist.py apply <xlsx>")
        apply(sys.argv[2])
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
